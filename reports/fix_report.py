from pathlib import Path
from io import BytesIO
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from PIL import Image

src=Path(r'C:\Users\ASUS\Downloads\Bao_cao_GitHub_Project_Cloudflare.docx')
out=Path(r'D:\IL\reports\Bao_cao_da_can_chinh.docx')
d=Document(src)
body=d.element.body
start=next(p._p for p in d.paragraphs if p.text=='8 Gửi thay đổi bằng pull request')
end=next(p._p for p in d.paragraphs if p.text=='9 Công khai website bằng Cloudflare')
children=list(body); a=children.index(start); b=children.index(end)
for el in children[a:b]: body.remove(el)

replacements={
'Git quản lý phiên bản tại máy tính. GitHub cung cấp repository, issue và pull request. GitHub Projects là bảng quản lý tiến độ. Trình soạn thảo dùng để sửa mã nguồn; trình duyệt dùng để kiểm thử; Cloudflare Workers phục vụ các tệp tĩnh của website.':
'Git quản lý phiên bản tại máy tính. GitHub lưu mã nguồn và các issue của nhóm. GitHub Projects theo dõi tiến độ. Trình soạn thảo dùng để sửa mã nguồn; trình duyệt dùng để kiểm thử; Cloudflare Workers phục vụ các tệp tĩnh của website.',
'Chuẩn bị tài khoản và Git → tạo repository → mời thành viên → lập Project và issue → phát triển trên nhánh → kiểm thử và hợp nhất → triển khai Cloudflare → kiểm tra website công khai.':
'Chuẩn bị tài khoản và Git → tạo repository → mời thành viên → lập Project và issue → cập nhật mã nguồn → kiểm thử → triển khai Cloudflare → kiểm tra website công khai.',
'Dự án có kho mã nguồn dùng chung, bảng quản lý công việc và website giới thiệu CV của 7 thành viên. Các thay đổi cần được lưu bằng commit để có thể xem lại lịch sử; issue và pull request giúp liên hệ yêu cầu với phần mã đã thực hiện.':
'Dự án có kho mã nguồn dùng chung, bảng quản lý công việc và website giới thiệu CV của 7 thành viên. Các thay đổi được lưu bằng commit để có thể xem lại lịch sử; issue giúp theo dõi yêu cầu và phần việc của nhóm.',
'Khi tiếp nhận thêm tài khoản cộng tác, người điều phối gửi lời mời repository, cấp quyền Project riêng và gán phần việc đã thống nhất. Mỗi lần thay đổi website, nhóm kiểm tra bản chạy trên máy, hợp nhất mã nguồn, triển khai và kiểm tra lại địa chỉ công khai.':
'Khi tiếp nhận thêm tài khoản cộng tác, người điều phối gửi lời mời repository, cấp quyền Project riêng và gán phần việc đã thống nhất. Mỗi lần thay đổi website, nhóm kiểm tra bản chạy trên máy, cập nhật mã nguồn, triển khai và kiểm tra lại địa chỉ công khai.'}
for p in d.paragraphs:
    old=p.text; new=replacements.get(old,old)
    if p.style.name.startswith('Heading'):
        m=re.match(r'^(9|10|11)(?= )',new)
        if m:new=str(int(m[1])-1)+new[len(m[1]):]
    if p.style.name=='Caption':
        new=new.replace('Hình 8.','Hình 7.').replace('Hình 9.','Hình 8.')
        captions={
            'Hình 4.':'Hình 4. Project hiển thị công việc theo trạng thái và người phụ trách',
            'Hình 5.':'Hình 5. Danh sách issue theo dõi công việc của nhóm',
            'Hình 7.':'Hình 7. Trang tổng quan Worker của dự án trên Cloudflare',
            'Hình 8.':'Hình 8. Trang chủ website công khai tại địa chỉ Cloudflare'
        }
        for prefix,value in captions.items():
            if new.startswith(prefix):new=value;break
    new=new.replace('triển khai bằng luồng ở mục 9.','triển khai bằng luồng ở mục 8.')
    if new!=old:p.text=new

for stname in ['Normal','Title','Heading 1','Heading 2','Caption']:
    st=d.styles[stname]; st.font.name='Times New Roman'; st.font.color.rgb=RGBColor(0,0,0)
    rf=st.element.get_or_add_rPr().rFonts
    for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),'Times New Roman')
    for k in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme','csTheme']:rf.attrib.pop(qn('w:'+k),None)
    for bd in list(st.element.get_or_add_pPr().findall(qn('w:pBdr'))):bd.getparent().remove(bd)
d.styles['Normal'].font.size=Pt(11.5)
d.styles['Normal'].paragraph_format.line_spacing=1.1
d.styles['Normal'].paragraph_format.space_after=Pt(5)
d.styles['Heading 1'].paragraph_format.page_break_before=True
d.styles['Heading 1'].paragraph_format.space_after=Pt(10)
d.styles['Heading 2'].paragraph_format.space_before=Pt(7)
d.styles['Heading 2'].paragraph_format.space_after=Pt(4)

# Replace fixed-height screenshot containers with inline images and their captions.
count=0
for t in list(d.tables):
    if len(t.rows)==1 and len(t.columns)==1:
        blips=t._tbl.xpath('.//a:blip')
        el=OxmlElement('w:p'); t._tbl.addprevious(el); p=Paragraph(el,d._body)
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next=True
        p.paragraph_format.keep_together=True
        p.paragraph_format.space_before=Pt(6)
        p.paragraph_format.space_after=Pt(4)
        p.paragraph_format.line_spacing=1
        for blip in blips:
            rid=blip.get(qn('r:embed')); data=d.part.related_parts[rid].blob
            im=Image.open(BytesIO(data)); w,h=im.size
            width=16.0; height=width*h/w
            if height>9.1:width*=9.1/height
            p.add_run().add_picture(BytesIO(data),width=Cm(width)); count+=1
        if not blips:
            p._element.getparent().remove(p._element)
        t._tbl.getparent().remove(t._tbl)

for p in list(d.paragraphs):
    if not p.text and not p._p.xpath('.//w:drawing'):
        # Empty page-break paragraphs from the original are superseded by heading pagination.
        if p._p.xpath('.//w:br[@w:type="page"]'):
            p._p.getparent().remove(p._p); continue
    if p.style.name=='Caption':
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_together=True
        p.paragraph_format.keep_with_next=False
        p.paragraph_format.space_after=Pt(5)
        for r in p.runs:r.font.size=Pt(10);r.italic=True
    elif p.style.name=='Normal' and not p._p.xpath('.//w:drawing'):
        p.paragraph_format.line_spacing=1.1
        p.paragraph_format.space_after=Pt(5)
    for r in p.runs:
        if r.font.name=='Times New Roman' and p.style.name=='Normal' and r.font.size and r.font.size.pt>=12:
            r.font.size=Pt(11.5)
for t in d.tables:
    for row in t.rows:
        trpr=row._tr.get_or_add_trPr()
        for x in list(trpr.findall(qn('w:trHeight'))):trpr.remove(x)
        if trpr.find(qn('w:cantSplit')) is None:trpr.append(OxmlElement('w:cantSplit'))
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
    t.autofit=False
d.save(out)
assert count==8,count
assert 'pull request' not in '\n'.join(p.text.lower() for p in d.paragraphs)
print(str(out)); print('Preserved screenshots:',count)
