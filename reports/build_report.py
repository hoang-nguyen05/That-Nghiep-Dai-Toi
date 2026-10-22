from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE

OUT=Path(r'D:\IL\reports')
OUT.mkdir(exist_ok=True)
d=Document()
s=d.sections[0]
s.page_height=Cm(29.7); s.page_width=Cm(21)
s.top_margin=Cm(2); s.bottom_margin=Cm(2)
s.left_margin=Cm(2.7); s.right_margin=Cm(2)
s.footer_distance=Cm(1)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    st=d.styles[name]; st.font.name='Times New Roman'; st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Times New Roman')
    for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme','csTheme']:
        st.element.rPr.rFonts.attrib.pop(qn('w:'+attr),None)
    pp=st.element.get_or_add_pPr()
    for border in list(pp.findall(qn('w:pBdr'))): pp.remove(border)
d.styles['Normal'].font.size=Pt(12)
d.styles['Normal'].paragraph_format.line_spacing=1.2
d.styles['Normal'].paragraph_format.space_after=Pt(7)
for name,size in [('Title',25),('Heading 1',17),('Heading 2',13)]:
    d.styles[name].font.size=Pt(size); d.styles[name].font.bold=True
d.styles['Heading 1'].paragraph_format.space_after=Pt(13)
d.styles['Heading 2'].paragraph_format.space_before=Pt(9)
d.styles['Caption'].font.size=Pt(11)
d.styles['Caption'].font.italic=True
f=s.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=f.add_run(); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); r._r.addnext(fld)
d.core_properties.title='Báo cáo quản lý dự án GitHub và triển khai website trên Cloudflare'
d.core_properties.author='Nhóm Thất Nghiệp Đại Tội'

def p(t='',bold=False):
    a=d.add_paragraph(); a.add_run(t).bold=bold; return a
def h(t): d.add_heading(t,2)
def page(t): d.add_page_break(); d.add_heading(t,1)
def steps(items):
    for i,t in enumerate(items,1): p(f'{i}. {t}')
def code(t):
    a=d.add_paragraph(); a.paragraph_format.line_spacing=1.05
    a.paragraph_format.space_after=Pt(9)
    r=a.add_run(t); r.font.name='Consolas'; r.font.size=Pt(9)
def table(headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,(a,w) in enumerate(zip(headers,widths)):
        t.columns[i].width=Cm(w); t.rows[0].cells[i].text=a
    for row in rows:
        c=t.add_row().cells
        for i,v in enumerate(row): c[i].text=v
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr(); cant=OxmlElement('w:cantSplit'); pr.append(cant)
        for i,c in enumerate(row.cells):
            c.width=Cm(widths[i]); tcpr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            tcpr.append(borders)
            for a in c.paragraphs:
                a.paragraph_format.space_after=Pt(4); a.paragraph_format.space_before=Pt(4); a.paragraph_format.line_spacing=1.05
                for r in a.runs: r.font.size=Pt(11); r.bold=ri==0
    return t
def shot(n,text,height=5.4):
    a=d.add_paragraph(); a.paragraph_format.space_after=Pt(2); a.paragraph_format.keep_with_next=True
    t=d.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.columns[0].width=Cm(16.3); row=t.rows[0]; row.height=Cm(height); row.height_rule=WD_ROW_HEIGHT_RULE.EXACTLY
    c=row.cells[0]; tcpr=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
    for edge in ['top','left','bottom','right']:
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'5'); e.set(qn('w:color'),'B8B8B8'); borders.append(e)
    tcpr.append(borders)
    a=d.add_paragraph(f'Hình {n}. {text}',style='Caption'); a.alignment=WD_ALIGN_PARAGRAPH.CENTER

# 1 Cover
for t in ['TRƯỜNG ......................................................................','KHOA ............................................................................']:
    a=p(t,True); a.alignment=WD_ALIGN_PARAGRAPH.CENTER
p();p()
a=d.add_paragraph('BÁO CÁO THỰC HÀNH',style='Title'); a.alignment=WD_ALIGN_PARAGRAPH.CENTER
a=d.add_paragraph('QUẢN LÝ DỰ ÁN TRÊN GITHUB\nVÀ TRIỂN KHAI WEBSITE\nTRÊN CLOUDFLARE',style='Title'); a.alignment=WD_ALIGN_PARAGRAPH.CENTER
p()
a=p('Dự án Thất Nghiệp Đại Tội',True); a.alignment=WD_ALIGN_PARAGRAPH.CENTER
a=p('Website giới thiệu CV của nhóm 7 thành viên'); a.alignment=WD_ALIGN_PARAGRAPH.CENTER
p();p()
for t in ['Học phần: ........................................................................................','Giảng viên hướng dẫn: ....................................................................','Lớp: ................................................................................................','Nhóm thực hiện: Thất Nghiệp Đại Tội','Người đại diện: Nguyễn Đắc Hoàng','Thời gian thực hiện: ........................................................................']: p(t)
p(); a=p('Năm 2026'); a.alignment=WD_ALIGN_PARAGRAPH.CENTER

# 2
page('1 Tổng quan đề tài')
p('Nhóm xây dựng website Thất Nghiệp Đại Tội để giới thiệu CV của 7 thành viên trên một trang chung. GitHub lưu trữ mã nguồn và lịch sử thay đổi; GitHub Projects theo dõi công việc; Cloudflare cung cấp địa chỉ website công khai để người xem truy cập bằng trình duyệt.')
h('1 1 Mục tiêu thực hiện')
p('Thiết lập kho mã nguồn dùng chung, hướng dẫn thành viên tham gia, tổ chức công việc theo trạng thái và đưa website HTML, CSS, JavaScript lên Internet. Sản phẩm cần mở được CV của từng người, hiển thị đủ hình ảnh và duy trì các hiệu ứng của giao diện hiện tại.')
h('1 2 Thành viên và nội dung website')
table(['STT','Họ và tên','Chủ đề đại tội'],[
['1','Nguyễn Đắc Hoàng','Dục Vọng'],['2','Hoàng Bắc','Kiêu Ngạo'],['3','Lường Thị Hạnh','Phàm Ăn'],['4','Dương Tuấn Hoàng','Phẫn Nộ'],['5','Nguyễn Huy Hoàng','Đố Kỵ'],['6','Nông Lâm Ngọc Sơn','Tham Lam'],['7','Vi Anh Quân','Lười Biếng']],[1.5,8.4,6.4])
h('1 3 Công cụ và phạm vi')
p('Git quản lý phiên bản tại máy tính. GitHub cung cấp repository, issue và pull request. GitHub Projects là bảng quản lý tiến độ. Trình soạn thảo dùng để sửa mã nguồn; trình duyệt dùng để kiểm thử; Cloudflare Workers phục vụ các tệp tĩnh của website.')
p('Website hiện dùng HTML, CSS và JavaScript, không có bước biên dịch ứng dụng. Nội dung CV được cập nhật trong tệp nguồn; thao tác chỉnh sửa trong bản xem không tự ghi dữ liệu về GitHub.')
h('1 4 Trình tự triển khai')
p('Chuẩn bị tài khoản và Git → tạo repository → mời thành viên → lập Project và issue → phát triển trên nhánh → kiểm thử và hợp nhất → triển khai Cloudflare → kiểm tra website công khai.')

#3
page('2 Chuẩn bị Git và tạo repository')
h('2 1 Chuẩn bị tài khoản và môi trường')
steps(['Tạo tài khoản GitHub tại github.com, xác minh email và đăng nhập. Mỗi thành viên dùng tài khoản riêng để lịch sử đóng góp có thể theo dõi được.','Cài Git từ git-scm.com nếu máy chưa có. Mở PowerShell hoặc Git Bash, chạy git --version để kiểm tra.','Thiết lập tên và email dùng cho commit. Thay nội dung ví dụ bằng thông tin của người thực hiện.'])
code('git --version\ngit config --global user.name "Ho va ten"\ngit config --global user.email "email@example.com"')
h('2 2 Tạo kho mã nguồn trên GitHub')
steps(['Chọn New repository trên GitHub; chọn chủ sở hữu là hoang-nguyen05.','Nhập tên That-Nghiep-Dai-Toi, thêm mô tả website CV của nhóm và chọn Public để công khai mã nguồn.','Nếu đưa một thư mục Git đã có lên lần đầu, tạo repository trống để tránh xung đột lịch sử. Nếu bắt đầu hoàn toàn trên GitHub, có thể tạo README rồi dùng git clone.','Nhấn Create repository và lưu địa chỉ HTTPS của repository. Hướng dẫn thao tác tham khảo tài liệu GitHub [1].'])
shot(1,'Màn hình tạo repository với tên dự án và chế độ Public',5.0)

#4
page('3 Đưa mã nguồn lên GitHub')
h('3 1 Khởi tạo từ thư mục có sẵn')
p('Với repository trống vừa tạo, mở terminal tại thư mục mã nguồn và thực hiện các lệnh dưới đây. Kiểm tra git status trước khi thêm tệp; không đưa mật khẩu, token hoặc tệp cấu hình chứa bí mật vào commit.')
code('git init\ngit branch -M main\ngit add .\ngit commit -m "Khoi tao website CV nhom"\ngit remote add origin https://github.com/hoang-nguyen05/That-Nghiep-Dai-Toi.git\ngit push -u origin main')
p('Đăng nhập GitHub qua trình quản lý thông tin xác thực khi được yêu cầu. Nếu thư mục đã là repository hoặc origin đã tồn tại, kiểm tra git status và git remote -v trước; không khởi tạo lại hoặc thêm trùng remote.')
h('3 2 Lấy dự án đã có về máy')
p('Với dự án hiện tại, thành viên mới dùng git clone thay cho quy trình khởi tạo ở trên. Sau khi sao chép, mở thư mục bằng trình soạn thảo để làm việc.')
code('git clone https://github.com/hoang-nguyen05/That-Nghiep-Dai-Toi.git\ncd That-Nghiep-Dai-Toi\ngit status')
h('3 3 Kiểm tra sau khi đẩy mã')
p('Mở repository trên trình duyệt, xác nhận nhánh main có index.html, style.css, script.js, assets và members. Lịch sử commit phải thể hiện nội dung thay đổi; README giới thiệu cách chạy dự án và đường dẫn website.')
shot(2,'Repository trên GitHub sau khi có mã nguồn và commit',5.0)

#5
page('4 Mời và tiếp nhận thành viên')
h('4 1 Thu thập thông tin tham gia nhóm')
p('Người điều phối thu thập họ tên, username GitHub, nội dung CV và phần việc mong muốn của từng người. Trước khi gán công việc, thống nhất thời hạn, tiêu chí hoàn thành và quy tắc gửi thay đổi. Username cần được kiểm tra đúng để tránh gửi lời mời nhầm tài khoản.')
h('4 2 Cấp quyền repository')
steps(['Chủ repository mở Settings → Collaborators trong nhóm mục Access.','Chọn Add people, nhập username hoặc email GitHub của thành viên và gửi lời mời.','Thành viên mở thông báo hoặc email để chấp nhận. Sau khi được chấp nhận, kiểm tra tài khoản đã có quyền cộng tác trong repository [2].','Thành viên clone repository về máy và thử tạo nhánh làm việc. Lặp lại quy trình cho các tài khoản còn lại.'])
h('4 3 Cấp quyền bảng Project riêng')
p('Quyền repository và quyền Project được quản lý riêng. Người điều phối mở Project → Settings → Manage access, thêm tài khoản với quyền Write để thành viên cập nhật công việc. Project công khai chỉ cho phép người ngoài xem; không tự cho mọi người quyền sửa bảng [4].')
p('Việc đã có CV trên website không đồng nghĩa tài khoản của người đó đã chấp nhận lời mời GitHub. Chỉ gán Assignees khi tài khoản thực tế đã tham gia; các vị trí chưa có tài khoản có thể để trống.')
shot(3,'Danh sách cộng tác viên hoặc lời mời tại Settings của repository',5.5)

#6
page('5 Tạo và liên kết GitHub Project')
h('5 1 Tạo bảng quản lý công việc')
steps(['Mở trang cá nhân GitHub → Projects → New project. Chọn bố cục Board và đặt tên dự án phù hợp [3].','Tạo hoặc kiểm tra ba trạng thái Todo, In progress và Done. Chọn khả năng hiển thị Public nếu muốn người xem theo dõi tiến độ công khai.','Mở repository → Projects để liên kết bảng với repository. Trong thiết lập Project, chọn repository làm kho liên kết nếu giao diện có mục tương ứng.','Mở Settings → Manage access và cấp quyền cho thành viên theo hướng dẫn ở mục 4.'])
h('5 2 Bảng Project của nhóm')
p('Dự án sử dụng Project số 1 thuộc tài khoản hoang-nguyen05. Mã nguồn nằm trong repository That-Nghiep-Dai-Toi; mỗi công việc có thể được tạo dưới dạng issue trong repository và thêm vào Project.')
p('Địa chỉ Project: https://github.com/users/hoang-nguyen05/projects/1')
h('5 3 Quy ước cập nhật tiến độ')
p('Todo là công việc đã xác định nhưng chưa bắt đầu. Khi nhận việc, thành viên chuyển sang In progress. Chỉ chuyển Done khi đáp ứng tiêu chí của issue, được kiểm tra và tích hợp vào bản chung; với việc liên quan website công khai, cần kiểm tra cả bản đã triển khai.')
shot(4,'Bảng Project với ba cột Todo, In progress và Done',6.0)

#7
page('6 Tạo issue và phân chia công việc')
p('Mỗi issue mô tả một phần việc có thể kiểm tra được. Nội dung gồm mục tiêu, tệp hoặc khu vực cần sửa, người phụ trách, thời hạn và tiêu chí hoàn thành. Dùng trường Projects để đưa issue vào bảng; dùng Assignees để chỉ định tài khoản thực hiện.')
h('6 1 Các gói công việc trên repository')
table(['Issue','Mã','Phần việc'],[
['#2','TV01','Điều phối, tích hợp và triển khai Cloudflare'],['#3','TV02','Nội dung và cấu trúc HTML'],['#4','TV03','Giao diện CSS trên máy tính'],['#5','TV04','Responsive và khả năng truy cập'],['#6','TV05','Tương tác JavaScript'],['#7','TV06','Kiểm thử và xuất PDF A4'],['#8','TV07','Tài liệu và báo cáo nhóm']],[1.3,1.7,13.3])
p('TV01–TV07 là mã phần việc, không mặc định tương ứng thứ tự họ tên ở mục 1. Người điều phối gán tài khoản phù hợp sau khi thống nhất trong nhóm.')
h('6 2 Ví dụ tiêu chí hoàn thành')
p('Với issue tương tác JavaScript: bấm tên hoặc ảnh thành viên phải mở đúng CV; nút đóng trả người xem về trang chính; liên kết mở tab mới hoạt động; các hiệu ứng cũ vẫn chạy và không xuất hiện lỗi JavaScript khi thao tác.')
shot(5,'Một issue có mô tả yêu cầu, Assignees và liên kết Project',5.5)

#8
page('7 Phát triển và kiểm tra trên máy')
h('7 1 Cấu trúc tệp chính')
table(['Tệp hoặc thư mục','Chức năng'],[
['index.html','Trang chủ và danh sách 7 thành viên'],['style.css và script.js','Giao diện, hiệu ứng và thao tác trên trang'],['assets/sins/','Bộ ảnh đại diện theo 7 đại tội'],['members/','Các trang CV và tài nguyên của từng thành viên'],['start.bat','Khởi chạy bản xem trên Windows'],['README.md và TEAM_GUIDE.md','Giới thiệu và quy trình cộng tác']],[5.3,11])
h('7 2 Chạy thử website')
p('Mở start.bat trong thư mục dự án. Khi máy có Python và cổng mặc định sẵn sàng, truy cập http://localhost:8080. Có thể mở terminal tại thư mục dự án để chạy trực tiếp lệnh sau; giữ terminal hoạt động trong lúc xem trang.')
code('python -m http.server 8080')
p('Nếu máy chưa có Python, cần cài Python hoặc dùng máy chủ tĩnh của trình soạn thảo. Kiểm tra tên tệp và đường dẫn bằng chữ hoa, chữ thường đúng như nguồn vì môi trường hosting có thể phân biệt chúng.')
h('7 3 Nội dung cần kiểm tra')
p('Thử toàn bộ thẻ thành viên, nút Xem CV, liên kết mở tab mới và thao tác đóng cửa sổ CV. Kiểm tra 7 ảnh đại diện, bố cục trên màn hình hẹp, âm thanh khi bật và các hiệu ứng ánh sáng, hạt nền, chuyển động khi rê chuột.')
shot(6,'Website chạy tại localhost và một CV được mở thành công',5.0)

#9
page('8 Gửi thay đổi bằng pull request')
h('8 1 Tạo nhánh và lưu thay đổi')
p('Trước khi sửa, đồng bộ nhánh main rồi tạo nhánh riêng theo nội dung công việc. Ví dụ dưới đây minh họa một nhánh sửa thao tác mở CV; thành viên thay tên nhánh và thông điệp commit cho phù hợp.')
code('git switch main\ngit pull origin main\ngit switch -c feature/mo-cv\n# Sua tep va chay thu website\ngit status\ngit add index.html script.js\ngit commit -m "Sua thao tac mo CV thanh vien"\ngit push -u origin feature/mo-cv')
h('8 2 Kiểm tra và hợp nhất')
steps(['Trên GitHub, tạo pull request từ nhánh vừa đẩy vào main. Mô tả vấn đề, thay đổi và cách đã kiểm tra.','Nếu pull request hoàn thành toàn bộ issue, thêm Closes #6 hoặc số issue tương ứng. Không đóng issue khi mới hoàn thành một phần.','Người kiểm tra đọc phần Files changed, mở thử website và xác nhận các tiêu chí. Nếu có xung đột, xử lý trên nhánh rồi kiểm tra lại.','Merge sau khi thay đổi đạt yêu cầu. Cập nhật trạng thái Project và triển khai bản đã hợp nhất lên Cloudflare.'])
h('8 3 Giữ bản chung hoạt động')
p('Chia nhỏ thay đổi để dễ xem xét. Khi sửa bố cục hoặc liên kết, kiểm tra cả chức năng CV và hiệu ứng hiện có. Không dùng git push --force trên main để thay thế lịch sử làm việc của nhóm.')
shot(7,'Pull request thể hiện thay đổi và kết quả hợp nhất vào main',5.2)

#10
page('9 Công khai website bằng Cloudflare')
h('9 1 Dịch vụ sử dụng')
p('Bản website của nhóm được triển khai bằng Cloudflare Workers với Static Assets. Các tệp HTML, CSS, JavaScript và hình ảnh được phục vụ trực tiếp qua HTTPS [5]. Địa chỉ thực tế dùng đuôi workers.dev; đây là cấu hình Workers của dự án, không phải địa chỉ Pages đuôi pages.dev.')
h('9 2 Chuẩn bị gói website')
p('Từ bản nguồn đã kiểm tra, tạo một thư mục phát hành chứa index.html, style.css, script.js, toàn bộ assets/ và members/. Nén nội dung thành ZIP sao cho index.html nằm ngay ở gốc sau khi giải nén. Loại thư mục .git, thông tin đăng nhập, tệp tạm và tài liệu không cần công khai.')
h('9 3 Tải lên và triển khai')
steps(['Đăng nhập Cloudflare, mở Workers & Pages. Với dự án đã có, chọn Worker that-nghiep-dai-toi. Với ứng dụng mới, chọn Create application và luồng tải tệp tĩnh phù hợp trên dashboard [6].','Chọn New deployment → Upload static files to update your Worker. Tải ZIP lên và kiểm tra danh sách tệp sau khi Cloudflare giải nén.','Nhấn Deploy, chờ thông báo thành công rồi mở địa chỉ production. Bản công khai của nhóm đã được cập nhật theo cách tải tệp tĩnh này.'])
p('Website production: https://that-nghiep-dai-toi.dtc235200306.workers.dev/')
shot(8,'Cloudflare nhận gói website và hiển thị triển khai thành công',5.5)

#11
page('10 Cập nhật và nghiệm thu website')
h('10 1 Liên kết repository để cập nhật')
p('Worker đã liên kết repository That-Nghiep-Dai-Toi, dùng nhánh production main, thư mục gốc / và lệnh deploy npx wrangler deploy. Website hiện không có build command. Khi dùng quy trình Git, cần kiểm tra log và deployment mới sau mỗi lần push; cấu hình Wrangler hoặc thiết lập tài nguyên tĩnh phải phù hợp với mã nguồn [5][6].')
p('Lần cập nhật đã thực hiện sử dụng tải ZIP trực tiếp. Việc push thành công trên GitHub chưa đủ chứng minh Cloudflare đã phục vụ bản mới. Nếu không có deployment mới từ Git, tải lại gói tệp đã kiểm tra và triển khai bằng luồng ở mục 9.')
h('10 2 Kết quả kiểm tra bản đã triển khai')
table(['Hạng mục','Kết quả ghi nhận'],[
['Trang chủ','Truy cập được qua địa chỉ HTTPS production'],['Ảnh đại tội','Cả 7 tệp ảnh tải thành công'],['Trang CV','Cả 7 đường dẫn CV trả về thành công'],['Mã JavaScript','Tệp phục vụ khớp bản nguồn sau sửa thao tác click']],[5.2,11.1])
h('10 3 Nghiệm thu thao tác trong trình duyệt')
p('Mở website công khai, làm mới trang và thử bấm tên Nguyễn Đắc Hoàng, ảnh đại diện, nút Xem CV, nút mở tab mới và nút đóng. Lặp lại với 6 thành viên còn lại. Kiểm tra màn hình điện thoại, hiệu ứng và đường dẫn ảnh; phản hồi HTTP thành công cần đi kèm thao tác giao diện đúng.')
shot(9,'Website công khai với thanh địa chỉ Cloudflare và CV đang mở',5.5)

#12
page('11 Đánh giá và tài liệu tham khảo')
h('11 1 Kết quả và hướng duy trì')
p('Dự án có kho mã nguồn dùng chung, bảng quản lý công việc và website giới thiệu CV của 7 thành viên. Các thay đổi cần được lưu bằng commit để có thể xem lại lịch sử; issue và pull request giúp liên hệ yêu cầu với phần mã đã thực hiện.')
p('Khi tiếp nhận thêm tài khoản cộng tác, người điều phối gửi lời mời repository, cấp quyền Project riêng và gán phần việc đã thống nhất. Mỗi lần thay đổi website, nhóm kiểm tra bản chạy trên máy, hợp nhất mã nguồn, triển khai và kiểm tra lại địa chỉ công khai.')
p('Phần cần tiếp tục xác nhận là cơ chế tự triển khai từ Git: chỉ xem là hoàn tất khi một commit mới tạo deployment thành công và nội dung production phản ánh đúng thay đổi. Cách tải ZIP trực tiếp hiện là quy trình cập nhật đã sử dụng được.')
h('11 2 Địa chỉ của dự án')
p('Repository: https://github.com/hoang-nguyen05/That-Nghiep-Dai-Toi')
p('Project: https://github.com/users/hoang-nguyen05/projects/1')
p('Website: https://that-nghiep-dai-toi.dtc235200306.workers.dev/')
h('11 3 Tài liệu tham khảo')
refs=[
('[1] GitHub Docs','Creating a new repository','https://docs.github.com/en/repositories/creating-and-managing-repositories/creating-a-new-repository'),
('[2] GitHub Docs','Inviting collaborators to a personal repository','https://docs.github.com/en/repositories/managing-your-repositorys-settings-and-features/repository-access-and-collaboration/inviting-collaborators-to-a-personal-repository'),
('[3] GitHub Docs','Creating a project','https://docs.github.com/en/issues/planning-and-tracking-with-projects/creating-projects/creating-a-project'),
('[4] GitHub Docs','Managing access to your projects','https://docs.github.com/en/issues/planning-and-tracking-with-projects/managing-your-project/managing-access-to-your-projects'),
('[5] Cloudflare Docs','Static Assets','https://developers.cloudflare.com/workers/static-assets/'),
('[6] Cloudflare Docs','Get started Dashboard','https://developers.cloudflare.com/workers/get-started/dashboard/')]
for org,title,url in refs:
    a=p(f'{org}. {title}.\n{url}')
    a.paragraph_format.line_spacing=1.0; a.paragraph_format.space_after=Pt(6)
    for r in a.runs:r.font.size=Pt(9)
p('Ngày tham khảo tài liệu: 21/10/2026.')
d.save(OUT/'Bao_cao_GitHub_Project_Cloudflare.docx')
print(OUT/'Bao_cao_GitHub_Project_Cloudflare.docx')
